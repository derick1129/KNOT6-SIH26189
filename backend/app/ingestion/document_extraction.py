"""
Multi-format evidence text/record extraction: PDF, DOCX, XLSX, JSON --
extending the ingestion pipeline beyond the original TXT/CSV support (see
app/api/routes/evidence.py's `_process()`, the one call site this module is
used from).

Every function here does real extraction with a real, maintained library
(pypdf, python-docx, openpyxl) -- never a stub that pretends to succeed.
When extraction genuinely can't produce usable content (a scanned/
image-only PDF with no text layer, a corrupt or password-protected file),
the caller gets an explicit failure reason back, not empty success: see
each function's `ExtractionResult`. OCR for scanned documents is
deliberately NOT implemented here -- this environment has no Tesseract (or
equivalent) binary installed (verified, not assumed), and the PS
instructions are explicit that a fake/stubbed OCR path is worse than
honestly reporting "requires OCR processing." A real OCR backend can be
added later as one more branch in `extract_text` without changing this
module's contract.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any


@dataclass
class TextExtractionResult:
    ok: bool
    text: str = ""
    warning: str = ""  # e.g. "Scanned document requires OCR processing." -- set only when ok is False


@dataclass
class RecordExtractionResult:
    ok: bool
    records: list[dict[str, Any]] = None  # type: ignore[assignment]
    warning: str = ""

    def __post_init__(self):
        if self.records is None:
            self.records = []


def extract_text_from_pdf(data: bytes) -> TextExtractionResult:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError:
        return TextExtractionResult(ok=False, warning="PDF support is not installed in this environment.")

    try:
        reader = PdfReader(BytesIO(data))
    except PdfReadError as exc:
        return TextExtractionResult(ok=False, warning=f"PDF could not be read (corrupt or encrypted): {exc}")

    if reader.is_encrypted:
        try:
            reader.decrypt("")  # only ever tries an empty password -- never guesses a real one
        except Exception:
            pass
        if reader.is_encrypted:
            return TextExtractionResult(ok=False, warning="PDF is password-protected; cannot extract text.")

    parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text)

    text = "\n\n".join(parts).strip()
    if not text:
        # No extractable text layer at all -- almost certainly a scanned/
        # image-only PDF. Reported honestly; no OCR is attempted (see
        # module docstring).
        return TextExtractionResult(ok=False, warning="Scanned document requires OCR processing.")
    return TextExtractionResult(ok=True, text=text)


def extract_text_from_docx(data: bytes) -> TextExtractionResult:
    try:
        import docx
    except ImportError:
        return TextExtractionResult(ok=False, warning="DOCX support is not installed in this environment.")

    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:  # python-docx raises plain Exception/PackageNotFoundError for bad files
        return TextExtractionResult(ok=False, warning=f"DOCX could not be read (corrupt or unsupported format): {exc}")

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    text = "\n".join(paragraphs).strip()
    if not text:
        return TextExtractionResult(ok=False, warning="DOCX contains no extractable text.")
    return TextExtractionResult(ok=True, text=text)


def extract_records_from_xlsx(data: bytes) -> RecordExtractionResult:
    """First sheet only, first row as headers -- the same CSV-shaped
    {column: value} record format app/ingestion/loaders.py's CDR/financial/
    criminal-history loaders already expect, so an XLSX upload reuses those
    loaders unchanged once converted to records here."""
    try:
        import openpyxl
    except ImportError:
        return RecordExtractionResult(ok=False, warning="XLSX support is not installed in this environment.")

    try:
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        return RecordExtractionResult(ok=False, warning=f"XLSX could not be read (corrupt or unsupported format): {exc}")

    sheet = wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return RecordExtractionResult(ok=False, warning="XLSX sheet is empty.")

    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header)]
    records = []
    for row in rows:
        if all(v is None for v in row):
            continue
        records.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})

    if not records:
        return RecordExtractionResult(ok=False, warning="XLSX has a header row but no data rows.")
    return RecordExtractionResult(ok=True, records=records)


def extract_json(data: bytes) -> tuple[TextExtractionResult | None, RecordExtractionResult | None]:
    """
    JSON can reasonably be either shape -- a list of structured records
    (CDR/financial/criminal-history-like) or a single document with a text
    field (matching the existing `/api/ingest/text` request body shape, or
    a bare string). Returns whichever extraction actually applies as a
    non-None result; the other is None, never both, so the caller doesn't
    have to guess which path succeeded.
    """
    import json as json_mod

    try:
        parsed = json_mod.loads(data.decode("utf-8-sig", errors="strict"))
    except (json_mod.JSONDecodeError, UnicodeDecodeError) as exc:
        return TextExtractionResult(ok=False, warning=f"JSON could not be parsed: {exc}"), None

    if isinstance(parsed, list) and parsed and all(isinstance(r, dict) for r in parsed):
        return None, RecordExtractionResult(ok=True, records=parsed)
    if isinstance(parsed, dict):
        text = parsed.get("text") or parsed.get("content") or parsed.get("body")
        if isinstance(text, str) and text.strip():
            return TextExtractionResult(ok=True, text=text), None
        return TextExtractionResult(ok=False, warning='JSON object has no "text"/"content"/"body" field.'), None
    if isinstance(parsed, str) and parsed.strip():
        return TextExtractionResult(ok=True, text=parsed), None

    return TextExtractionResult(ok=False, warning="JSON structure not recognized (expected a list of records, "
                                                    "an object with a text field, or a bare string)."), None
