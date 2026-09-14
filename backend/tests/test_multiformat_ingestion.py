"""
Multi-format evidence ingestion tests: PDF, DOCX, XLSX, JSON, and honest
unsupported/scanned handling -- exercised through the real upload endpoint
with real files built by the same libraries doing the extraction (pypdf,
python-docx, openpyxl), not mocked.
"""
from __future__ import annotations

import io
import json

from tests.conftest import auth


def _make_investigation_and_case(client, token, name):
    inv = client.post("/api/investigations", json={"name": name}, headers=auth(token)).json()
    case = client.post(f"/api/investigations/{inv['id']}/cases",
                        json={"case_number": f"MF-{inv['id'][:8]}", "title": "Multi-format test case"},
                        headers=auth(token)).json()
    return inv, case


def _upload(client, token, inv, case, filename: str, content: bytes, source_type: str, mime: str = "application/octet-stream"):
    return client.post(
        f"/api/investigations/{inv['id']}/cases/{case['id']}/evidence",
        headers=auth(token), data={"source_type": source_type},
        files={"file": (filename, io.BytesIO(content), mime)},
    )


def _make_real_pdf(text: str) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    # pypdf has no built-in "write text to a page" helper (that's reportlab's
    # job) -- build a minimal one-page PDF with a real content stream by
    # hand, using pypdf's own low-level objects, so the round-trip through
    # pypdf's OWN reader (extract_text_from_pdf) is genuine, not mocked.
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = DecodedStreamObject()
    content = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    stream.set_data(content)
    page = writer.pages[0]
    resources = DictionaryObject()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    resources[NameObject("/Font")] = DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = writer._add_object(stream)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_real_docx(text: str) -> bytes:
    import docx
    doc = docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_real_xlsx(headers: list[str], rows: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_with_real_text_is_processed(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "PDF Text")
    pdf_bytes = _make_real_pdf("Arjun Malhotra called Priya Nair near Sector 5.")
    res = _upload(client, investigator_token, inv, case, "report.pdf", pdf_bytes, "FIR", "application/pdf")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "PROCESSED"
    assert body["processing_summary"]["entities_extracted"] >= 1


def test_scanned_pdf_with_no_text_layer_is_honestly_reported(client, investigator_token):
    """A PDF with a page but no text content stream at all -- pypdf's
    extract_text() returns "" for it, exactly like a scanned/image-only
    PDF would, exercising the real "no extractable text" path rather than
    mocking it."""
    from pypdf import PdfWriter
    inv, case = _make_investigation_and_case(client, investigator_token, "PDF Scanned")
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    res = _upload(client, investigator_token, inv, case, "scanned.pdf", buf.getvalue(), "FIR", "application/pdf")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "FAILED"
    assert "OCR" in body["processing_summary"]["error"]


def test_corrupt_pdf_is_honestly_reported(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "PDF Corrupt")
    res = _upload(client, investigator_token, inv, case, "bad.pdf", b"not a real pdf file", "FIR", "application/pdf")
    body = res.json()
    assert body["processing_status"] == "FAILED"
    assert body["processing_summary"]["error"]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_with_real_text_is_processed(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "DOCX Text")
    docx_bytes = _make_real_docx("Rohit Bhatia transferred funds to Neha Iyer.")
    res = _upload(client, investigator_token, inv, case, "intel.docx", docx_bytes, "INTELLIGENCE_REPORT",
                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "PROCESSED"
    assert body["processing_summary"]["entities_extracted"] >= 1


def test_corrupt_docx_is_honestly_reported(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "DOCX Corrupt")
    res = _upload(client, investigator_token, inv, case, "bad.docx", b"not a real docx", "FIR")
    body = res.json()
    assert body["processing_status"] == "FAILED"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def test_xlsx_structured_records_processed(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "XLSX Structured")
    xlsx_bytes = _make_real_xlsx(
        ["from_account", "to_account", "amount", "timestamp"],
        [["XLACC1", "XLACC2", 50000, "2026-01-01T10:00:00"]],
    )
    res = _upload(client, investigator_token, inv, case, "transactions.xlsx", xlsx_bytes, "FINANCIAL_TRANSACTION",
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "PROCESSED"
    assert body["processing_summary"]["entities_created"] >= 2

    graph = client.get("/api/graph", params={"investigation_id": inv["id"]}, headers=auth(investigator_token)).json()
    labels = {n["label"] for n in graph["nodes"]}
    assert "XLACC1" in labels and "XLACC2" in labels


def test_xlsx_empty_sheet_is_honestly_reported(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "XLSX Empty")
    xlsx_bytes = _make_real_xlsx(["from_account", "to_account", "amount"], [])
    res = _upload(client, investigator_token, inv, case, "empty.xlsx", xlsx_bytes, "FINANCIAL_TRANSACTION")
    body = res.json()
    assert body["processing_status"] == "FAILED"
    assert "no data rows" in body["processing_summary"]["error"]


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------

def test_json_text_document_processed(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "JSON Text")
    payload = json.dumps({"text": "Kavita Rao met Sameer Joshi at Central Station."}).encode()
    res = _upload(client, investigator_token, inv, case, "report.json", payload, "FIR", "application/json")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "PROCESSED"
    assert body["processing_summary"]["entities_extracted"] >= 1


def test_json_structured_records_processed(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "JSON Structured")
    payload = json.dumps([{"from_account": "JACC1", "to_account": "JACC2", "amount": 1000,
                            "timestamp": "2026-01-01T10:00:00"}]).encode()
    res = _upload(client, investigator_token, inv, case, "transactions.json", payload, "FINANCIAL_TRANSACTION")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "PROCESSED"


def test_malformed_json_is_honestly_reported(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "JSON Malformed")
    res = _upload(client, investigator_token, inv, case, "bad.json", b"{not valid json", "FIR")
    body = res.json()
    assert body["processing_status"] == "FAILED"


# ---------------------------------------------------------------------------
# Unsupported formats
# ---------------------------------------------------------------------------

def test_unsupported_extension_is_honestly_reported_not_silently_succeeded(client, investigator_token):
    inv, case = _make_investigation_and_case(client, investigator_token, "Unsupported Format")
    res = _upload(client, investigator_token, inv, case, "report.exe", b"binary junk", "FIR")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["processing_status"] == "FAILED"
    assert "Unsupported file type" in body["processing_summary"]["error"]
